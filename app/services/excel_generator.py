"""Excel generation service for approved onboarding documents."""

import logging
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from app.core.config import Settings, get_settings

logger = logging.getLogger(__name__)

HEADER_FIELDS = (
    "employee_code",
    "candidate_name",
    "father_name",
    "date_of_birth",
    "date_of_joining",
    "aadhaar_number",
    "pan_number",
    "gender",
    "marital_status",
    "address",
    "bank_account_number",
    "ifsc_code",
)

HEADER_LABELS = {
    "employee_code": "Employee Code",
    "candidate_name": "Candidate Name",
    "father_name": "Father's Name",
    "date_of_birth": "Date of Birth",
    "date_of_joining": "Date of Joining",
    "aadhaar_number": "Aadhaar Number",
    "pan_number": "PAN Number",
    "gender": "Gender",
    "marital_status": "Marital Status",
    "address": "Address",
    "bank_account_number": "Bank Account Number",
    "ifsc_code": "IFSC Code",
}


class ExcelGenerationError(RuntimeError):
    """Raised when Excel file generation fails."""


def generate_onboarding_excel(
    documents: list[dict[str, Any]],
    batch_id: str,
    settings: Settings | None = None,
) -> str:
    """Generate an Excel workbook from approved onboarding documents.

    Returns the file path of the generated Excel file.
    """
    runtime_settings = settings or get_settings()

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Onboarding Data"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = ["#", "Filename"] + [HEADER_LABELS.get(f, f) for f in HEADER_FIELDS]
        ws.append(headers)
        ws.row_dimensions[1].height = 30

        for col_idx, _header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, doc in enumerate(documents, 2):
            data = doc.get("verified_json") or doc.get("extracted_json") or {}
            filename = doc.get("filename", "")
            row_data = [row_idx - 1, filename]

            for field in HEADER_FIELDS:
                field_data = data.get(field, {})
                value = field_data.get("value") if isinstance(field_data, dict) else field_data
                row_data.append(value if value is not None else "")

            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 30
        for col_idx in range(3, len(headers) + 1):
            ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else f"A{col_idx - 1}"].width = 22

        output_dir = Path(runtime_settings.upload_temp_dir) / batch_id
        output_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        output_path = output_dir / f"onboarding_data_{timestamp}.xlsx"
        wb.save(str(output_path))

        logger.info(
            "Onboarding Excel generated",
            extra={
                "event": "excel_generated",
                "batch_id": batch_id,
                "file_path": str(output_path),
                "document_count": len(documents),
            },
        )

        return str(output_path)
    except Exception as exc:
        logger.exception(
            "Failed to generate onboarding Excel",
            extra={"event": "excel_generation_failed", "batch_id": batch_id},
        )
        raise ExcelGenerationError(f"Failed to generate Excel file: {exc}") from exc
