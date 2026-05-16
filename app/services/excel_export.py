"""Excel export service for approved onboarding documents."""

import logging
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from app.core.config import Settings, get_settings

logger = logging.getLogger(__name__)

EXPORT_FIELDS = (
    "employee_code",
    "candidate_name",
    "father_name",
    "date_of_birth",
    "date_of_joining",
    "aadhaar_number",
    "pan_number",
    "gender",
    "marital_status",
    "address",
    "bank_account_number",
    "ifsc_code",
)

HEADER_LABELS: dict[str, str] = {
    "employee_code": "Employee Code",
    "candidate_name": "Candidate Name",
    "father_name": "Father's Name",
    "date_of_birth": "Date of Birth",
    "date_of_joining": "Date of Joining",
    "aadhaar_number": "Aadhaar Number",
    "pan_number": "PAN Number",
    "gender": "Gender",
    "marital_status": "Marital Status",
    "address": "Address",
    "bank_account_number": "Bank Account Number",
    "ifsc_code": "IFSC Code",
}


class ExcelExportError(RuntimeError):
    """Raised when Excel export generation fails."""


def export_approved_documents(
    documents: list[dict[str, Any]],
    settings: Settings | None = None,
) -> str:
    """Generate an Excel workbook with one row per approved document.

    Includes all extracted fields: employee_code, candidate_name,
    father_name, date_of_birth, date_of_joining, aadhaar_number,
    pan_number, gender, marital_status, address, bank_account_number,
    and ifsc_code.

    Returns the absolute file path of the generated Excel file.
    """
    runtime_settings = settings or get_settings()

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Approved Documents"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = ["#", "Filename"] + [HEADER_LABELS[f] for f in EXPORT_FIELDS]
        ws.append(headers)
        ws.row_dimensions[1].height = 30

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, doc in enumerate(documents, 2):
            data = doc.get("verified_json") or doc.get("extracted_json") or {}
            filename = doc.get("filename", "")
            row_data: list[Any] = [row_idx - 1, filename]

            for field in EXPORT_FIELDS:
                field_data = data.get(field, {})
                if isinstance(field_data, dict):
                    value = field_data.get("value")
                else:
                    value = field_data
                row_data.append(value if value is not None else "")

            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 30
        for col_idx in range(3, len(headers) + 1):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = 24

        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        filename = f"approved_documents_{timestamp}.xlsx"

        output_dir = Path(runtime_settings.upload_temp_dir) / "exports"
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / filename
        wb.save(str(output_path))

        logger.info(
            "Approved documents Excel exported",
            extra={
                "event": "excel_exported",
                "file_path": str(output_path),
                "document_count": len(documents),
            },
        )

        return str(output_path)
    except Exception as exc:
        logger.exception(
            "Failed to export approved documents Excel",
            extra={"event": "excel_export_failed"},
        )
        raise ExcelExportError(f"Failed to export Excel: {exc}") from exc
